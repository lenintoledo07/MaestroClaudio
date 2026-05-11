"""Google Drive: metadata, descarga streaming, listado de carpetas y extractores
locales para PDF/PPTX. Reutiliza el refresh-token-handling de auth_service.
"""

from __future__ import annotations

import io
import logging
import os
import time
from typing import Iterable
from uuid import UUID

import asyncpg
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseDownload

from config import settings
from services.auth_service import get_valid_google_token

logger = logging.getLogger(__name__)

# Mime → tipo interno
_MIME_MAP = {
    "video/mp4": "video",
    "video/quicktime": "video",
    "video/x-matroska": "video",
    "video/x-msvideo": "video",
    "audio/mpeg": "video",  # tratamos audios como "video" para el pipeline
    "audio/mp4": "video",
    "audio/x-m4a": "video",
    "audio/wav": "video",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.ms-powerpoint": "pptx",
}


async def _credentials_for_user(user_id: UUID) -> Credentials:
    """Construye Credentials con un access_token siempre válido."""
    conn = await asyncpg.connect(settings.DATABASE_URL, command_timeout=10)
    try:
        token = await get_valid_google_token(conn, user_id)
    finally:
        await conn.close()
    return Credentials(token=token["access_token"])


def detect_file_type(mime_type: str) -> str:
    """Convierte el mimeType de Drive a uno de nuestros tipos internos."""
    return _MIME_MAP.get(mime_type, "unknown")


async def get_file_metadata(user_id: UUID, file_id: str) -> dict:
    """Devuelve {id, name, mimeType, size, webViewLink, type, duration_ms?}."""
    creds = await _credentials_for_user(user_id)
    service = build("drive", "v3", credentials=creds, cache_discovery=False)
    try:
        f = service.files().get(
            fileId=file_id,
            fields="id, name, mimeType, size, webViewLink, videoMediaMetadata",
            supportsAllDrives=True,
        ).execute()
    except HttpError as exc:
        if exc.resp.status == 404:
            raise FileNotFoundError(f"Archivo {file_id} no encontrado en Drive") from exc
        if exc.resp.status == 403:
            raise PermissionError(f"Sin permisos para {file_id}") from exc
        raise

    duration_ms = None
    vmm = f.get("videoMediaMetadata") or {}
    if vmm.get("durationMillis"):
        duration_ms = int(vmm["durationMillis"])

    return {
        "id": f["id"],
        "name": f.get("name", file_id),
        "mimeType": f.get("mimeType", ""),
        "size": int(f["size"]) if f.get("size") else None,
        "webViewLink": f.get("webViewLink"),
        "type": detect_file_type(f.get("mimeType", "")),
        "duration_ms": duration_ms,
    }


async def download_file_temporarily(
    user_id: UUID, file_id: str, dest_dir: str = "/tmp"
) -> str:
    """Descarga el archivo a {dest_dir}/{file_id}.{ext} con streaming.

    Retorna el path local. El caller es responsable de borrarlo después.
    """
    meta = await get_file_metadata(user_id, file_id)
    ext = _ext_for_mime(meta["mimeType"]) or "bin"
    local_path = os.path.join(dest_dir, f"{file_id}.{ext}")

    creds = await _credentials_for_user(user_id)
    service = build("drive", "v3", credentials=creds, cache_discovery=False)

    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)

    started = time.perf_counter()
    with open(local_path, "wb") as fh:
        downloader = MediaIoBaseDownload(fh, request, chunksize=8 * 1024 * 1024)
        done = False
        while not done:
            _, done = downloader.next_chunk(num_retries=3)

    elapsed = time.perf_counter() - started
    size_mb = os.path.getsize(local_path) / (1024 * 1024)
    logger.info(
        "Drive download %s → %s (%.1f MB en %.1fs)",
        file_id, local_path, size_mb, elapsed,
    )
    return local_path


async def list_folder_contents(user_id: UUID, folder_id: str) -> list[dict]:
    """Lista archivos y subcarpetas dentro de una carpeta de Drive."""
    creds = await _credentials_for_user(user_id)
    service = build("drive", "v3", credentials=creds, cache_discovery=False)

    items: list[dict] = []
    page_token = None
    while True:
        resp = service.files().list(
            q=f"'{folder_id}' in parents and trashed = false",
            fields="nextPageToken, files(id, name, mimeType, size, modifiedTime)",
            pageSize=1000,
            pageToken=page_token,
            supportsAllDrives=True,
            includeItemsFromAllDrives=True,
        ).execute()
        items.extend(resp.get("files", []))
        page_token = resp.get("nextPageToken")
        if not page_token:
            break
    return items


async def list_subfolders(user_id: UUID, folder_id: str) -> list[dict]:
    """Lista SOLO subcarpetas (no archivos). Útil para que el user elija
    qué carpeta de su Drive raíz vincular a cada materia."""
    creds = await _credentials_for_user(user_id)
    service = build("drive", "v3", credentials=creds, cache_discovery=False)

    items: list[dict] = []
    page_token = None
    while True:
        resp = service.files().list(
            q=(
                f"'{folder_id}' in parents "
                f"and mimeType = 'application/vnd.google-apps.folder' "
                f"and trashed = false"
            ),
            fields="nextPageToken, files(id, name, modifiedTime)",
            orderBy="name",
            pageSize=200,
            pageToken=page_token,
            supportsAllDrives=True,
            includeItemsFromAllDrives=True,
        ).execute()
        items.extend(resp.get("files", []))
        page_token = resp.get("nextPageToken")
        if not page_token:
            break
    return items


async def get_folder_metadata(user_id: UUID, folder_id: str) -> dict:
    """Devuelve nombre + id de una carpeta. Sirve para verificar que
    el ID que pegó el user es realmente una carpeta accesible."""
    creds = await _credentials_for_user(user_id)
    service = build("drive", "v3", credentials=creds, cache_discovery=False)
    f = service.files().get(
        fileId=folder_id,
        fields="id, name, mimeType",
        supportsAllDrives=True,
    ).execute()
    if f.get("mimeType") != "application/vnd.google-apps.folder":
        raise ValueError(f"El ID {folder_id} no es una carpeta de Drive")
    return {"id": f["id"], "name": f["name"]}


def extract_file_id_from_url(url: str) -> str | None:
    """Extrae el file_id de una URL de Drive (varios formatos)."""
    import re

    patterns = [
        r"/file/d/([a-zA-Z0-9_-]+)",        # /file/d/{ID}/view
        r"[?&]id=([a-zA-Z0-9_-]+)",          # ?id={ID}
        r"/folders/([a-zA-Z0-9_-]+)",        # /folders/{ID}
    ]
    for p in patterns:
        m = re.search(p, url)
        if m:
            return m.group(1)
    return None


# ── Extractores locales (PDF, PPTX) ─────────────────────────────────────────


def extract_text_pdf(local_path: str) -> str:
    """Extrae texto de un PDF con PyMuPDF."""
    import fitz  # PyMuPDF

    parts: list[str] = []
    with fitz.open(local_path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    return "\n\n".join(p for p in parts if p.strip())


def extract_text_pptx(local_path: str) -> str:
    """Extrae texto de un PPTX (incluye notas del orador)."""
    from pptx import Presentation

    prs = Presentation(local_path)
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_text: list[str] = [f"[Slide {i}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                slide_text.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_text.append(f"[Notas] {notes}")
        parts.append("\n".join(t for t in slide_text if t.strip()))
    return "\n\n".join(parts)


# ── Helpers ──────────────────────────────────────────────────────────────────


def _ext_for_mime(mime: str) -> str | None:
    table = {
        "video/mp4": "mp4",
        "video/quicktime": "mov",
        "video/x-matroska": "mkv",
        "audio/mpeg": "mp3",
        "audio/mp4": "m4a",
        "audio/x-m4a": "m4a",
        "audio/wav": "wav",
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
        "application/vnd.ms-powerpoint": "ppt",
    }
    return table.get(mime)


# Suprimimos export simbólico de Iterable / io para evitar warnings
_ = Iterable, io
